"""用大模型起草 / 审查分类标准：从标准文档或已标注样本生成标注指南草稿，并做“标准体检”。"""
from __future__ import annotations

import asyncio
import copy
import io
import json
import random
import re
from pathlib import Path

import httpx

from .config import Config, ModelConfig, TaskConfig, LabelDef
from .llm import LLMError, create_llm
from .prompts import render_guide

MAX_DOC_CHARS = 40000
SYSTEM = "你是资深的数据标注规范设计专家，擅长把业务分类标准整理成标注员可以直接执行的标注指南。你只输出 JSON。"

_OUTPUT_SPEC = (
    '输出格式（只输出这个 JSON）：\n'
    '{"name": "任务名称", "description": "一句话任务说明：分类对象是什么、每条只归入一个类别",\n'
    ' "labels": [{"name": "类别名", "definition": "判断依据", "examples": ["正例"], "counter_examples": ["反例"]}],\n'
    ' "rules": ["边界规则"], "notes": ["需要人工确认的问题"]}'
)


def read_document(name: str, data: bytes) -> tuple[str, bool]:
    """把 txt / md / docx / pdf 转为纯文本，返回 (文本, 是否被截断)。"""
    suffix = Path(name).suffix.lower()
    if suffix == ".docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)
    elif suffix == ".pdf":
        import fitz

        with fitz.open(stream=data, filetype="pdf") as doc:
            text = "\n".join(page.get_text() for page in doc)
    elif suffix in (".txt", ".md", ".markdown", ".csv"):
        for enc in ("utf-8-sig", "gbk"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError("无法识别文本编码，请另存为 UTF-8")
    else:
        raise ValueError("支持 .docx / .pdf / .txt / .md 文件")
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise ValueError("没有从文件中读到文字（扫描版 PDF 需要先做文字识别）")
    return text[:MAX_DOC_CHARS], len(text) > MAX_DOC_CHARS


def _task_block(task: dict) -> str:
    return render_guide(TaskConfig(
        name=task.get("name", ""), description=task.get("description", ""),
        labels=[LabelDef(lab["name"], lab.get("definition", ""), list(lab.get("examples") or []),
                         list(lab.get("counter_examples") or [])) for lab in task.get("labels") or []],
        rules=list(task.get("rules") or []),
    ))


def prompt_from_document(doc: str, current: dict | None = None) -> str:
    base = ("\n6. 下面是当前已有的标准，请在它的基础上按文档修订：保持已有类别名称不变，只修改与文档不符或缺失的内容。\n"
            f"<current>\n{_task_block(current)}\n</current>\n") if current and current.get("labels") else ""
    return (
        "下面是一份分类标准文档（可能是制度文件、标注说明或业务口径）。请把它整理成可直接用于大模型分类的标注指南。\n\n"
        "要求：\n"
        "1. 类别名称沿用文档中的叫法，不要改名、合并或拆分；文档没有的类别不要编造。文档覆盖不到的情况如果很常见，在 notes 中建议是否需要“其他”类。\n"
        "2. definition：一两句话写清判断依据，忠实于文档，不要加入文档没有的口径。\n"
        "3. examples：每类 2~4 条简短、典型的正例，文档里有例子优先用文档的；counter_examples：看起来像但不属于该类的情况，没有就留空。\n"
        "4. rules：边界规则，写清楚容易混淆的两类怎么区分、同时符合多个类别时怎么判，每条一句话。\n"
        "5. notes：文档中含糊、互相矛盾或遗漏的地方，每条一句话，供人工确认。"
        f"{base}\n\n{_OUTPUT_SPEC}\n\n# 分类标准文档\n<document>\n{doc}\n</document>"
    )


def sample_examples(pairs: list[tuple[str, str]], per_label: int = 25, max_chars: int = 150, seed: int = 42) -> list[tuple[str, str]]:
    by: dict[str, list[str]] = {}
    for text, label in pairs:
        by.setdefault(label, []).append(text)
    rng = random.Random(seed)
    out = []
    for label, texts in by.items():
        picked = rng.sample(texts, min(per_label, len(texts)))
        out += [(" ".join(t.split())[:max_chars], label) for t in picked]
    return out


def prompt_from_examples(samples: list[tuple[str, str]], description: str = "") -> str:
    labels = list(dict.fromkeys(label for _, label in samples))
    lines = "\n".join(f"【{label}】{text}" for text, label in samples)
    desc = f"\n任务背景：{description}\n" if description else ""
    return (
        "下面是一批已由人工标注的样本（每类抽样若干条）。请归纳这批数据的标注标准，整理成标注指南。\n\n"
        "要求：\n"
        f"1. 类别名称必须严格使用这些：{' / '.join(labels)}，不增不减、不改名。\n"
        "2. definition：根据样本归纳这一类的共同特征和判断依据，尤其要写出这批数据特有的标注习惯（例如某些看似中性的表述被标成了某一类）。\n"
        "3. examples：每类选 2~4 条最典型的样本原文（可截短）；counter_examples：选容易被误判成这一类、实际标为其他类的样本。\n"
        "4. rules：归纳类与类之间的边界，每条一句话，要具体到可以执行。\n"
        "5. notes：样本中看起来标注不一致、或难以归纳的地方。"
        f"{desc}\n\n{_OUTPUT_SPEC}\n\n# 已标注样本\n{lines}"
    )


def prompt_review(task: dict) -> str:
    return (
        "请审查下面这份分类标注指南，找出会导致标注员（或模型）判断不一致的问题：\n"
        "- 重叠：两个类别的定义有交叉，同一条文本可能同时符合；\n"
        "- 模糊：定义用词含糊，缺少可操作的判断依据；\n"
        "- 缺失：常见情况没有被任何类别覆盖，或缺少关键的边界规则；\n"
        "- 冲突：规则之间、规则与定义之间互相矛盾；\n"
        "- 示例问题：正例 / 反例与定义不符。\n"
        "只指出真正会影响判断的问题，不要凑数；每个问题给出具体的修改建议（可以直接给出改写后的定义或新增的规则）。\n\n"
        '输出格式（只输出这个 JSON）：\n{"issues": [{"type": "重叠|模糊|缺失|冲突|示例问题", "labels": ["涉及的类别"], '
        '"problem": "问题", "suggestion": "修改建议"}], "summary": "一句话总体评价"}\n\n'
        f"# 标注指南\n{_task_block(task)}"
    )


def _extract_json(text: str) -> dict:
    text = re.sub(r"<think>.*?</think>", "", text, flags=re.S)
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end <= start:
        raise LLMError("模型回复中没有 JSON")
    try:
        obj = json.loads(text[start:end + 1])
    except json.JSONDecodeError as e:
        raise LLMError(f"模型回复的 JSON 不完整（可能超出输出长度）：{e}") from e
    if not isinstance(obj, dict):
        raise LLMError("模型回复的 JSON 不是对象")
    return obj


def _strs(v) -> list[str]:
    return [str(x).strip() for x in (v or []) if str(x).strip()] if isinstance(v, list) else []


def normalize_draft(obj: dict) -> dict:
    labels = []
    for lab in obj.get("labels") or []:
        if not isinstance(lab, dict) or not str(lab.get("name", "")).strip():
            continue
        d = {"name": str(lab["name"]).strip(), "definition": str(lab.get("definition", "")).strip()}
        for key in ("examples", "counter_examples"):
            if _strs(lab.get(key)):
                d[key] = _strs(lab.get(key))
        labels.append(d)
    if len(labels) < 2:
        raise LLMError("草稿中的类别少于 2 个，请检查文档内容或换一个模型重试")
    return {"name": str(obj.get("name", "")).strip(), "description": str(obj.get("description", "")).strip(),
            "labels": labels, "rules": _strs(obj.get("rules")), "notes": _strs(obj.get("notes"))}


def drafting_models(config: Config) -> list[ModelConfig]:
    """可用于起草的模型：启用的大模型投票模型和仲裁模型（本地小模型、mock 不行）。"""
    ms = config.models + ([config.arbiter] if config.arbiter else [])
    return [m for m in ms if m.provider in ("openai", "anthropic")]


async def _ask(config: Config, model: ModelConfig, user: str, max_tokens: int) -> str:
    cfg = copy.deepcopy(model)
    cfg.max_tokens = max(cfg.max_tokens, max_tokens)
    cfg.logprobs = False
    cfg.json_mode = False
    async with httpx.AsyncClient(timeout=max(config.pipeline.timeout, 240)) as http:
        return (await create_llm(cfg, http, config).complete(SYSTEM, user)).text


def ask_json(config: Config, model: ModelConfig, user: str, max_tokens: int = 6000) -> dict:
    return _extract_json(asyncio.run(_ask(config, model, user, max_tokens)))


def draft_from_document(config: Config, model: ModelConfig, doc: str, current: dict | None = None) -> dict:
    return normalize_draft(ask_json(config, model, prompt_from_document(doc, current)))


def draft_from_examples(config: Config, model: ModelConfig, pairs: list[tuple[str, str]], description: str = "",
                        per_label: int = 25) -> dict:
    draft = normalize_draft(ask_json(config, model, prompt_from_examples(sample_examples(pairs, per_label), description)))
    wanted = list(dict.fromkeys(label for _, label in pairs))
    got = {lab["name"] for lab in draft["labels"]}
    if got != set(wanted):
        draft["notes"].insert(0, f"模型给出的类别 {sorted(got)} 与数据中的类别 {wanted} 不一致，已按数据中的类别修正")
        by = {lab["name"]: lab for lab in draft["labels"]}
        draft["labels"] = [by.get(name, {"name": name, "definition": ""}) for name in wanted]
    return draft


def review_criteria(config: Config, model: ModelConfig, task: dict) -> dict:
    obj = ask_json(config, model, prompt_review(task), max_tokens=4000)
    issues = [i for i in obj.get("issues") or [] if isinstance(i, dict) and i.get("problem")]
    return {"issues": issues, "summary": str(obj.get("summary", "")).strip()}
